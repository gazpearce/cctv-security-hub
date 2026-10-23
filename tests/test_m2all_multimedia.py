#!/usr/bin/env python3
"""
tests/test_m2all_multimedia.py
==============================
Comprehensive Automated Test Suite for Milestone M2-ALL:
1. FAQ Metadata Catalog Validity (34+ FAQs across 6 service pillars, Rule 30 hook architecture, word counts, image paths)
2. 16:9 PPTX and Vector PDF Slide Deck Generation (Dimensions, 4-slide structure, NumberedCanvas, E-E-A-T credentials)
3. Binary EXIF Geocoding Injection & Extraction (Pillow + piexif, GPS coordinates 53.7489° N, 1.6628° W, Latin-1 safe IFD0)
4. RSS 2.0 XML Schema & Enclosure Integrity (podcast.xml, faq_multimedia_feed.xml, iTunes namespace, CDATA unescaping)
5. 14-Column Layout Invariant in doc3 (Rows 341-360 repair, zero 9-col rows, Doc3RegistryGuard prevention)
6. Rule 45 Zero Naked Root Link Compliance across all catalog targets and ledger entries
7. Browser UI Worker FAQ Upload Hook Verification

Author: Gary Pearce / ADI Communications
Repository: github-cctv-seo
"""

import os
import re
import csv
import json
import unittest
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET

import piexif
from PIL import Image
from pptx import Presentation

# Project imports
from produce_faq_multimedia import (
    REPO_ROOT,
    CATALOG_PATH,
    DOC3_CSV_PATH,
    PODCAST_FEED_PATH,
    FAQ_FEED_PATH,
    DEFAULT_LAT,
    DEFAULT_LON,
    GOOGLE_SLIDES_APEX_URL,
    YOUTUBE_APEX_URL,
    BLOGGER_SHIELD_URL,
    Doc3RegistryGuard,
    FAQSlideGenerator,
    FAQExifImageGenerator,
    RSSFeedGenerator,
    FAQVideoGenerator,
    deg_to_dms_rational,
    audit_multimedia_assets,
)
from browser_ui_worker import BrowserUIWorker, AssetUploadManifest


class TestFaqMetadataCatalog(unittest.TestCase):
    """Verifies that the catalog contains at least 34 high-intent FAQs across all 6 service pillars."""

    @classmethod
    def setUpClass(cls):
        cls.assertTrue(CATALOG_PATH.exists(), f"Catalog file missing at {CATALOG_PATH}")
        with open(CATALOG_PATH, "r", encoding="utf-8") as f:
            cls.catalog = json.load(f)
        cls.faqs = cls.catalog.get("faqs", [])

    def test_01_total_faq_count_and_version(self):
        self.assertGreaterEqual(len(self.faqs), 34, f"Expected at least 34 FAQs, found {len(self.faqs)}")
        self.assertIn("version", self.catalog)

    def test_02_all_six_service_pillars_represented(self):
        expected_services = {
            "CCTV Engineering",
            "Satellite Telecoms",
            "TV & Media Walls",
            "Smart Alarms",
            "TV & Broadcast",
            "Structured Cabling"
        }
        actual_services = {f.get("service") for f in self.faqs}
        self.assertEqual(expected_services, actual_services, f"Missing services: {expected_services - actual_services}")

        # Each service must have at least 5 FAQs
        service_counts = {}
        for f in self.faqs:
            s = f.get("service")
            service_counts[s] = service_counts.get(s, 0) + 1
        for s, count in service_counts.items():
            self.assertGreaterEqual(count, 5, f"Service '{s}' has only {count} FAQs (minimum 5 required)")

    def test_03_faq_item_mandatory_fields(self):
        required_fields = [
            "faq_id", "service", "question", "title", "slug",
            "matching_image_asset", "technical_standards", "buffer_url",
            "blogger_shield_url", "script_narration", "word_count", "tags", "timestamps"
        ]
        for f in self.faqs:
            fid = f.get("faq_id", "UNKNOWN")
            for field in required_fields:
                self.assertIn(field, f, f"FAQ {fid} missing field '{field}'")
                self.assertTrue(f[field], f"FAQ {fid} field '{field}' is empty")

    def test_04_matching_images_exist_on_disk(self):
        for f in self.faqs:
            rel_img = f.get("matching_image_asset")
            img_path = REPO_ROOT / rel_img
            self.assertTrue(img_path.exists(), f"Matching image does not exist on disk for {f.get('faq_id')}: {img_path}")

    def test_05_rule_30_hook_and_credentials_architecture(self):
        """Verifies that script narration opens with a question/hook, not a self-ad, and contains Gary's phone."""
        for f in self.faqs:
            fid = f.get("faq_id")
            script = f.get("script_narration", "")
            # Must NOT open with "I'm Gary Pearce" or "Hello, I'm"
            first_sentence = script.split(".")[0].strip()
            self.assertFalse(first_sentence.startswith("I'm Gary Pearce"), f"FAQ {fid} violates Rule 30 by starting with identity: '{first_sentence}'")
            self.assertFalse(first_sentence.startswith("My name is Gary"), f"FAQ {fid} violates Rule 30 by starting with identity: '{first_sentence}'")
            # Must contain direct mobile: 07860 244 274
            self.assertIn("07860 244 274", script, f"FAQ {fid} missing direct mobile 07860 244 274")
            # Word count should be 70 to 110 words
            words = len(script.split())
            self.assertTrue(70 <= words <= 110, f"FAQ {fid} word count {words} out of expected range 70-110")


class TestSlideGeneration(unittest.TestCase):
    """Verifies that 16:9 PPTX and vector PDF slide decks are generated with full fidelity."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.out_dir = Path(self.temp_dir.name)
        self.generator = FAQSlideGenerator(output_dir=self.out_dir, docs_dir=self.out_dir)

        with open(CATALOG_PATH, "r", encoding="utf-8") as f:
            cat = json.load(f)
        self.faq_sample = cat["faqs"][0]

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_01_generate_pptx_dimensions_and_structure(self):
        pptx_path = self.generator.generate_pptx(self.faq_sample)
        self.assertTrue(pptx_path.exists(), f"PPTX not generated at {pptx_path}")
        self.assertGreater(pptx_path.stat().st_size, 15000, "PPTX file size abnormally small")

        prs = Presentation(str(pptx_path))
        # 4 slides
        self.assertEqual(len(prs.slides), 4, f"Expected exactly 4 slides, got {len(prs.slides)}")
        # 16:9 widescreen layout
        self.assertAlmostEqual(prs.slide_width.inches, 13.333, places=2)
        self.assertAlmostEqual(prs.slide_height.inches, 7.500, places=2)

        # Slide 1 text checks
        slide1_text = "".join([shape.text_frame.text for shape in prs.slides[0].shapes if shape.has_text_frame])
        self.assertIn("BOTTOM LINE UP FRONT", slide1_text)
        self.assertIn("07860 244 274", slide1_text)
        self.assertIn("9.98/10", slide1_text)

        # Slide 4 text checks (Credentials & Apex)
        slide4_text = "".join([shape.text_frame.text for shape in prs.slides[3].shapes if shape.has_text_frame])
        self.assertIn("GARY PEARCE", slide4_text)
        self.assertIn(GOOGLE_SLIDES_APEX_URL, slide4_text)

    def test_02_generate_vector_pdf_with_numbered_canvas(self):
        import base64, zlib
        pdf_path = self.generator.generate_pdf(self.faq_sample)
        self.assertTrue(pdf_path.exists(), f"Vector PDF not generated at {pdf_path}")
        self.assertGreater(pdf_path.stat().st_size, 3000, "PDF file size abnormally small")

        # Read PDF binary and decompress Ascii85 + Flate stream to verify headers/footers
        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()

        self.assertTrue(pdf_bytes.startswith(b"%PDF-"), "File is not a valid PDF")
        idx1 = pdf_bytes.find(b"stream") + len(b"stream")
        idx2 = pdf_bytes.find(b"endstream")
        stream_raw = pdf_bytes[idx1:idx2].strip()
        if not stream_raw.startswith(b"<~"):
            stream_raw = b"<~" + stream_raw
        decompressed = zlib.decompress(base64.a85decode(stream_raw, adobe=True))

        self.assertIn(b"ADI COMMUNICATIONS", decompressed)
        self.assertIn(b"07860 244 274", decompressed)
        self.assertIn(b"Rule 45", decompressed)


class TestExifGeocoding(unittest.TestCase):
    """Verifies that binary EXIF GPS metadata is correctly injected and readable."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.out_dir = Path(self.temp_dir.name)
        self.generator = FAQExifImageGenerator(output_dir=self.out_dir)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_01_deg_to_dms_rational(self):
        dms = deg_to_dms_rational(53.7489)
        self.assertEqual(dms[0], (53, 1))
        self.assertEqual(dms[1], (44, 1))
        self.assertEqual(dms[2], (560400, 10000))

        dms_lon = deg_to_dms_rational(-1.6628)
        self.assertEqual(dms_lon[0], (1, 1))
        self.assertEqual(dms_lon[1], (39, 1))
        self.assertEqual(dms_lon[2], (460800, 10000))

    def test_02_exif_injection_and_readback(self):
        meta = {
            "title": "Hikvision ColorVu vs Infrared CCTV Test",
            "artist": "Gary Pearce - ADI Communications (07860 244 274)",
            "copyright": "Copyright 2026 ADI Communications Ltd. Gary Pearce 07860 244 274",
            "lat": DEFAULT_LAT,
            "lon": DEFAULT_LON,
            "user_comment": "Gary Pearce 07860 244 274 | CCTV | Lat 53.7489 Lon -1.6628 | Rule 45"
        }
        webp_path, jpg_path = self.generator.process_and_tag(None, "test_asset", meta)

        self.assertTrue(webp_path.exists())
        self.assertTrue(jpg_path.exists())

        # Inspect JPEG EXIF
        jpg_loaded = piexif.load(str(jpg_path))
        artist = jpg_loaded["0th"][piexif.ImageIFD.Artist].decode("latin-1")
        self.assertIn("Gary Pearce", artist)
        self.assertIn("07860 244 274", artist)

        gps_lat = jpg_loaded["GPS"][piexif.GPSIFD.GPSLatitude]
        gps_lon = jpg_loaded["GPS"][piexif.GPSIFD.GPSLongitude]
        lat_ref = jpg_loaded["GPS"][piexif.GPSIFD.GPSLatitudeRef].decode("latin-1")
        lon_ref = jpg_loaded["GPS"][piexif.GPSIFD.GPSLongitudeRef].decode("latin-1")

        self.assertEqual(lat_ref, "N")
        self.assertEqual(lon_ref, "W")
        self.assertEqual(gps_lat[0], (53, 1))
        self.assertEqual(gps_lon[0], (1, 1))


class TestRSSFeedStructure(unittest.TestCase):
    """Verifies that podcast.xml and faq_multimedia_feed.xml adhere to RSS 2.0 standards."""

    def test_01_podcast_xml_exists_and_parses(self):
        self.assertTrue(PODCAST_FEED_PATH.exists())
        tree = ET.parse(str(PODCAST_FEED_PATH))
        root = tree.getroot()
        self.assertEqual(root.tag, "rss")
        self.assertEqual(root.attrib.get("version"), "2.0")

        channel = root.find("channel")
        self.assertIsNotNone(channel)
        items = channel.findall("item")
        self.assertGreaterEqual(len(items), 3, "Expected at least 3 master podcast episodes")

        for it in items:
            self.assertIsNotNone(it.find("title"))
            self.assertIsNotNone(it.find("guid"))
            enc = it.find("enclosure")
            self.assertIsNotNone(enc)
            self.assertTrue(enc.attrib.get("url"))
            self.assertTrue(int(enc.attrib.get("length", 0)) > 0)
            self.assertEqual(enc.attrib.get("type"), "audio/mpeg")

    def test_02_faq_feed_xml_item_injection_idempotency(self):
        self.assertTrue(FAQ_FEED_PATH.exists())
        with open(CATALOG_PATH, "r", encoding="utf-8") as f:
            cat = json.load(f)
        faq = cat["faqs"][0]

        item_el = RSSFeedGenerator.build_item_element(
            faq,
            REPO_ROOT / faq["matching_image_asset"],
            "https://gary-pearce-forensic-hub.web.app/assets/videos/test.mp4",
            "https://gary-pearce-forensic-hub.web.app/assets/videos/test_thumb.webp",
            35.0,
            GOOGLE_SLIDES_APEX_URL,
            faq["buffer_url"]
        )

        initial_tree = ET.parse(str(FAQ_FEED_PATH))
        initial_count = len(initial_tree.getroot().findall("./channel/item"))

        # Append item
        appended = RSSFeedGenerator.append_to_feed(FAQ_FEED_PATH, item_el)
        self.assertTrue(appended)

        after_tree = ET.parse(str(FAQ_FEED_PATH))
        new_count = len(after_tree.getroot().findall("./channel/item"))
        self.assertTrue(new_count >= initial_count)

        # Idempotency check: re-appending the exact same item must not increment count
        re_append = RSSFeedGenerator.append_to_feed(FAQ_FEED_PATH, item_el)
        self.assertTrue(re_append)
        after_re_tree = ET.parse(str(FAQ_FEED_PATH))
        re_count = len(after_re_tree.getroot().findall("./channel/item"))
        self.assertEqual(new_count, re_count, "RSS append_to_feed created duplicate item GUID")


class TestDoc3RegistryCompliance(unittest.TestCase):
    """Verifies that doc3_published_asset_links.csv maintains 100% 14-column layout."""

    def test_01_canonical_14_column_header(self):
        self.assertTrue(DOC3_CSV_PATH.exists())
        with open(DOC3_CSV_PATH, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            header = next(reader)
        self.assertEqual(len(header), 14, f"Header has {len(header)} columns instead of 14")
        expected_header = [
            "asset_id", "parent_profile_id", "platform_name", "profile_live_url",
            "asset_title", "asset_type", "asset_public_url", "top_of_stack_apex_url",
            "shield_target_url", "seo_tags", "seo_description_summary",
            "dofollow_status", "publication_date", "verification_status"
        ]
        self.assertEqual(header, expected_header)

    def test_02_all_rows_strictly_14_columns(self):
        with open(DOC3_CSV_PATH, "r", encoding="utf-8") as f:
            reader = list(csv.reader(f))

        malformed = [(i + 1, len(r)) for i, r in enumerate(reader) if len(r) != 14]
        self.assertEqual(len(malformed), 0, f"Found {len(malformed)} rows not having 14 columns: {malformed[:5]}")

    def test_03_repaired_rows_341_to_360(self):
        with open(DOC3_CSV_PATH, "r", encoding="utf-8") as f:
            reader = list(csv.reader(f))

        self.assertEqual(len(reader), 360, f"Expected 360 rows in doc3, got {len(reader)}")
        for line_num in range(341, 361):
            row = reader[line_num - 1]
            self.assertEqual(len(row), 14, f"Line {line_num} does not have 14 columns")
            expected_id = f"LINK-{line_num:04d}"
            self.assertEqual(row[0], expected_id, f"Line {line_num} asset_id is '{row[0]}', expected '{expected_id}'")
            self.assertEqual(row[1], "PROF-IMG-02")
            self.assertEqual(row[2], "Catbox CDN")
            self.assertEqual(row[5], "CDN Image Asset")
            self.assertTrue(row[6].startswith("https://files.catbox.moe/"))
            self.assertEqual(row[7], GOOGLE_SLIDES_APEX_URL)
            self.assertTrue(row[8].startswith("https://adicommunications.co.uk/cctv-installation-"))
            self.assertEqual(row[13], "HTTP_200")

    def test_04_doc3_registry_guard_enforcement(self):
        next_id = Doc3RegistryGuard.get_next_asset_id(DOC3_CSV_PATH)
        self.assertTrue(next_id.startswith("LINK-"))
        num = int(next_id.split("-")[1])
        self.assertGreaterEqual(num, 361)

        # Rule 45 enforcement: Attempting to register a naked root link must raise ValueError
        bad_entry = {
            "platform_name": "Test Platform",
            "shield_target_url": "https://adicommunications.co.uk/",
            "asset_public_url": "https://example.com/asset.mp4"
        }
        with self.assertRaises(ValueError):
            Doc3RegistryGuard.register_asset(DOC3_CSV_PATH, bad_entry)


class TestRule45Compliance(unittest.TestCase):
    """Verifies that no generated assets or catalog links contain naked root links."""

    def test_01_catalog_urls_rule_45_clean(self):
        with open(CATALOG_PATH, "r", encoding="utf-8") as f:
            cat = json.load(f)

        for f in cat.get("faqs", []):
            buf = f.get("buffer_url", "")
            self.assertFalse(
                re.search(r"^https?://(www\.)?adicommunications\.co\.uk/?$", buf.strip(), re.I),
                f"FAQ {f.get('faq_id')} has raw root link: {buf}"
            )
            # Ensure it is a deep path (not empty or just /)
            parsed = re.sub(r"^https?://[^/]+", "", buf)
            self.assertTrue(
                len(parsed.strip("/")) > 3,
                f"FAQ {f.get('faq_id')} buffer_url is not a deep service/buffer page: {buf}"
            )

    def test_02_doc3_shield_targets_rule_45_clean(self):
        with open(DOC3_CSV_PATH, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                shield_url = row.get("shield_target_url", "")
                if shield_url:
                    self.assertFalse(
                        re.search(r"^https?://(www\.)?adicommunications\.co\.uk/?$", shield_url.strip(), re.I),
                        f"doc3 asset {row.get('asset_id')} violates Rule 45 with raw root link: {shield_url}"
                    )


class TestBrowserUIWorkerFAQHooks(unittest.TestCase):
    """Verifies that browser_ui_worker.py has operational hooks for manual FAQ uploading."""

    def test_01_worker_has_upload_faq_asset_method(self):
        worker = BrowserUIWorker()
        self.assertTrue(hasattr(worker, "upload_faq_asset"))
        self.assertTrue(callable(getattr(worker, "upload_faq_asset")))

    def test_02_upload_faq_asset_validation(self):
        import asyncio
        worker = BrowserUIWorker()
        platform_dummy = {"name": "Test Host", "signup_url": "https://example.com"}

        # Non-existent FAQ ID should return failure dictionary gracefully
        res = asyncio.run(worker.upload_faq_asset("FAQ-NON-EXISTENT-999", "video", platform_dummy))
        self.assertEqual(res.get("status"), "FAILED")
        self.assertIn("not found in catalog", res.get("error", ""))


class TestSystemAuditCLI(unittest.TestCase):
    """Verifies that produce_faq_multimedia.py --audit passes cleanly."""

    def test_01_system_audit_pass(self):
        res = audit_multimedia_assets()
        self.assertEqual(res.get("status"), "PASS")
        self.assertEqual(len(res.get("missing_images", [])), 0)
        self.assertTrue(res.get("doc3_valid"))


if __name__ == "__main__":
    unittest.main()
